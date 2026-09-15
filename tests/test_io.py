from __future__ import annotations

import io
import json

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from sito.core.io import (
    export_filename,
    format_data_value,
    guess_mapping,
    import_rows,
    keyword_table,
    parse_data_value,
    read_table,
    snapshot_table,
    write_table,
)
from sito.core.snapshots import create_snapshot
from sito.models import Keyword, Project
from sito.plugins.base import PluginError

# --- read_table --------------------------------------------------------------------------


def test_read_csv_comma():
    data = b"Keyword,Volume\nbuy server,100\ncheap vps,50\nrent dedicated,30\n"
    table = read_table("keywords.csv", data)
    assert table.columns == ["Keyword", "Volume"]
    assert table.rows[0] == {"Keyword": "buy server", "Volume": "100"}
    assert len(table.rows) == 3


def test_read_semicolon_delimited():
    data = b"Keyword;Volume\nbuy server;100\ncheap vps;50\nrent dedicated;30\n"
    table = read_table("keywords.csv", data)
    assert table.columns == ["Keyword", "Volume"]
    assert table.rows[0]["Volume"] == "100"


def test_read_utf16_tab_separated_planner_export():
    text = "Keyword\tAvg. monthly searches\nbuy server\t100\ncheap vps\t50\nrent dedicated\t30\n"
    data = b"\xff\xfe" + text.encode("utf-16-le")
    table = read_table("planner.csv", data)
    assert table.columns == ["Keyword", "Avg. monthly searches"]
    assert table.rows[0]["Keyword"] == "buy server"
    assert table.rows[0]["Avg. monthly searches"] == "100"


def test_read_cp1251():
    text = "Ключевое слово,Частотность\nкупить сервер,100\nдешевый vps,50\n"
    data = text.encode("cp1251")
    table = read_table("keywords.csv", data)
    assert table.rows[0]["Ключевое слово"] == "купить сервер"
    assert table.rows[0]["Частотность"] == "100"


def test_read_txt_one_keyword_per_line():
    data = b"buy server\ncheap vps\n\nrent dedicated\n"
    table = read_table("list.txt", data)
    assert table.columns == ["Keyword"]
    assert [r["Keyword"] for r in table.rows] == ["buy server", "cheap vps", "rent dedicated"]


def test_read_txt_drops_recognized_header():
    data = "Запрос\nкупить сервер\nдешевый vps\n".encode()
    table = read_table("list.txt", data)
    assert table.columns == ["Keyword"]
    assert [r["Keyword"] for r in table.rows] == ["купить сервер", "дешевый vps"]


def test_read_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.append(["Keyword", "Volume"])
    ws.append(["buy server", 100])
    ws.append(["cheap vps", 50])
    buf = io.BytesIO()
    wb.save(buf)
    table = read_table("keywords.xlsx", buf.getvalue())
    assert table.columns == ["Keyword", "Volume"]
    assert table.rows[0] == {"Keyword": "buy server", "Volume": "100"}


def test_read_table_unsupported_extension():
    with pytest.raises(PluginError):
        read_table("data.pdf", b"whatever")


# --- guess_mapping -------------------------------------------------------------------------


def test_guess_mapping_ahrefs():
    mapping = guess_mapping(["Keyword", "Volume", "KD", "CPC", "Parent Topic"])
    assert mapping["Keyword"] == "keyword"
    assert mapping["Volume"] == "volume"
    assert mapping["KD"] == "kd"
    assert mapping["CPC"] == "cpc"
    assert mapping["Parent Topic"] == "data:parent_topic"


def test_guess_mapping_semrush():
    mapping = guess_mapping(["Keyword", "Search Volume", "Keyword Difficulty", "CPC"])
    assert mapping["Search Volume"] == "volume"
    assert mapping["Keyword Difficulty"] == "kd"


def test_guess_mapping_keyword_planner():
    mapping = guess_mapping(["Keyword", "Avg. monthly searches"])
    assert mapping["Keyword"] == "keyword"
    assert mapping["Avg. monthly searches"] == "volume"


def test_guess_mapping_russian_and_ukrainian():
    mapping = guess_mapping(["Ключевое слово", "Частотность", "Запрос", "Обсяг", "Ключове слово"])
    assert mapping["Ключевое слово"] == "keyword"
    assert mapping["Частотность"] == "volume"
    assert mapping["Запрос"] == "keyword"
    assert mapping["Обсяг"] == "volume"
    assert mapping["Ключове слово"] == "keyword"


def test_guess_mapping_round_trips_sito_export():
    columns = ["keyword", "volume", "kd", "cpc", "source", "cluster", "status", "reason", "relevance"]
    mapping = guess_mapping(columns)
    assert mapping["keyword"] == "keyword"
    assert mapping["volume"] == "volume"
    assert mapping["cluster"] == "ignore"
    assert mapping["status"] == "excluded"
    assert mapping["reason"] == "excluded_reason"
    assert mapping["relevance"] == "data:relevance"


def test_guess_mapping_reverses_fixed_column_collision():
    # keyword_table() escapes a data key equal to a fixed column as "data.<key>" on export;
    # guess_mapping() must reverse that literally, not slugify the dotted header.
    mapping = guess_mapping(["keyword", "data.volume", "data.status"])
    assert mapping["data.volume"] == "data:volume"
    assert mapping["data.status"] == "data:status"


# --- parse_data_value / format_data_value ---------------------------------------------------


def test_parse_data_value_scalars():
    assert parse_data_value("8") == 8
    assert parse_data_value("8.5") == 8.5
    assert parse_data_value("true") is True
    assert parse_data_value("False") is False
    assert parse_data_value("null") is None
    assert parse_data_value("none") is None
    assert parse_data_value("commercial") == "commercial"
    assert parse_data_value('["a", "b"]') == ["a", "b"]
    assert parse_data_value('{"h1": "Buy"}') == {"h1": "Buy"}
    assert parse_data_value("[not json") == "[not json"  # invalid JSON stays text


def test_format_data_value_matches_parse_data_value():
    for value in (8, 8.5, True, False, None, "commercial", ["a", "b"], {"h1": "Buy"}):
        assert parse_data_value(format_data_value(value)) == value
    assert format_data_value("plain text") == "plain text"  # strings are not JSON-quoted


# --- import_rows ---------------------------------------------------------------------------


def test_import_rows_add_mode(session_factory, project):
    rows = [{"Keyword": "buy server", "Volume": "100"}, {"Keyword": "cheap vps", "Volume": "50"}]
    mapping = {"Keyword": "keyword", "Volume": "volume"}
    with session_factory() as s:
        result = import_rows(s, project.id, rows, mapping, "add", "test")
        s.commit()
        assert result.added == 2
        assert result.skipped == 0
        kws = {k.keyword: k for k in s.scalars(select(Keyword).where(Keyword.project_id == project.id))}
    assert kws["buy server"].volume == 100


def test_import_rows_add_mode_skips_existing(session_factory, project):
    with session_factory() as s:
        s.add(Keyword(project_id=project.id, keyword="buy server", source="manual"))
        s.commit()
    rows = [{"Keyword": "buy server"}, {"Keyword": "new term"}]
    mapping = {"Keyword": "keyword"}
    with session_factory() as s:
        result = import_rows(s, project.id, rows, mapping, "add", "test")
        s.commit()
    assert result.added == 1
    assert result.skipped == 1


def test_import_rows_update_mode_only_touches_existing(session_factory, project):
    with session_factory() as s:
        s.add(Keyword(project_id=project.id, keyword="buy server", source="manual", volume=10))
        s.commit()
    rows = [{"Keyword": "buy server", "Volume": "200"}, {"Keyword": "new term", "Volume": "5"}]
    mapping = {"Keyword": "keyword", "Volume": "volume"}
    with session_factory() as s:
        result = import_rows(s, project.id, rows, mapping, "update", "test")
        s.commit()
    assert result.updated == 1
    assert result.skipped == 1
    with session_factory() as s:
        kw = s.scalars(select(Keyword).where(Keyword.project_id == project.id, Keyword.keyword == "buy server")).one()
        assert kw.volume == 200
        assert s.scalar(select(Keyword).where(Keyword.keyword == "new term")) is None


def test_import_rows_merge_mode_empty_values_do_not_overwrite(session_factory, project):
    with session_factory() as s:
        s.add(Keyword(project_id=project.id, keyword="buy server", source="manual", volume=10, kd=5.0))
        s.commit()
    rows = [{"Keyword": "buy server", "Volume": "", "KD": "8"}]
    mapping = {"Keyword": "keyword", "Volume": "volume", "KD": "kd"}
    with session_factory() as s:
        result = import_rows(s, project.id, rows, mapping, "merge", "test")
        s.commit()
    assert result.updated == 1
    with session_factory() as s:
        kw = s.scalars(select(Keyword).where(Keyword.project_id == project.id, Keyword.keyword == "buy server")).one()
        assert kw.volume == 10
        assert kw.kd == 8.0


def test_import_rows_status_and_reason_round_trip(session_factory, project):
    rows = [{"Keyword": "buy server", "status": "retained", "reason": "too broad"}]
    mapping = {"Keyword": "keyword", "status": "excluded", "reason": "excluded_reason"}
    with session_factory() as s:
        import_rows(s, project.id, rows, mapping, "add", "test")
        s.commit()
        kw = s.scalars(select(Keyword).where(Keyword.project_id == project.id)).one()
        assert kw.excluded is True
        assert kw.excluded_reason == "too broad"


def test_import_rows_data_field_and_json_value(session_factory, project):
    rows = [{"Keyword": "buy server", "intent": "commercial", "tags": '["a", "b"]'}]
    mapping = {"Keyword": "keyword", "intent": "data:intent", "tags": "data:tags"}
    with session_factory() as s:
        import_rows(s, project.id, rows, mapping, "add", "test")
        s.commit()
        kw = s.scalars(select(Keyword).where(Keyword.project_id == project.id)).one()
        assert kw.data["intent"] == "commercial"
        assert kw.data["tags"] == ["a", "b"]


# --- export ----------------------------------------------------------------------------------


def test_keyword_table_includes_data_columns(session_factory, project):
    with session_factory() as s:
        s.add(Keyword(project_id=project.id, keyword="buy server", source="manual", volume=100, data={"relevance": 8}))
        s.commit()
    with session_factory() as s:
        columns, rows = keyword_table(s, project.id, "all")
    assert "relevance" in columns
    assert rows[0]["relevance"] == 8
    assert rows[0]["status"] == "passing"


def test_write_table_csv_round_trip(session_factory, project):
    with session_factory() as s:
        s.add(Keyword(project_id=project.id, keyword="buy server", source="manual", volume=100))
        s.commit()
    with session_factory() as s:
        columns, rows = keyword_table(s, project.id, "all")
    payload = write_table(columns, rows, "csv")
    text = payload.decode("utf-8-sig")
    assert "keyword" in text.splitlines()[0]
    assert "buy server" in text


def test_write_table_xlsx_round_trip(session_factory, project):
    with session_factory() as s:
        s.add(Keyword(project_id=project.id, keyword="buy server", source="manual", data={"tags": ["a", "b"]}))
        s.commit()
    with session_factory() as s:
        columns, rows = keyword_table(s, project.id, "all")
    payload = write_table(columns, rows, "xlsx")
    wb = load_workbook(io.BytesIO(payload))
    ws = wb.active
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == columns
    body = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
    assert "buy server" in body


def test_write_table_json_round_trip(session_factory, project):
    with session_factory() as s:
        s.add(Keyword(project_id=project.id, keyword="buy server", source="manual", data={"tags": ["a", "b"]}))
        s.commit()
    with session_factory() as s:
        columns, rows = keyword_table(s, project.id, "all")
    payload = write_table(columns, rows, "json")
    parsed = json.loads(payload)
    assert parsed[0]["keyword"] == "buy server"
    assert parsed[0]["tags"] == ["a", "b"]  # JSON export keeps native lists/dicts (unlike csv/xlsx)


def test_export_filename_is_ascii_slug_and_timestamped():
    name = export_filename("My Project", "passing keywords", "xlsx")
    assert name.endswith(".xlsx")
    assert " " not in name
    assert name.startswith("my-project-passing-keywords-")


def test_snapshot_table_reads_frozen_payload(session_factory, project):
    with session_factory() as s:
        s.add(Keyword(project_id=project.id, keyword="buy server", source="manual", volume=100))
        s.commit()
        snap = create_snapshot(s, project.id, "checkpoint")
        s.commit()
        columns, rows = snapshot_table(snap, "all")
    assert "keyword" in columns
    assert rows[0]["keyword"] == "buy server"


# --- real export -> re-import round trips ---------------------------------------------------

ORIGINAL_DATA = {
    "relevance": 8,
    "score": 3.5,
    "is_branded": False,
    "tags": ["a", "b"],
    "notes": "plain text",
    "volume": 999,  # collides with the fixed "volume" column - exported as "data.volume"
}


def _reimport_into_fresh_project(session_factory, columns, rows, fmt):
    """Write `rows` as `fmt`, then reimport (mode "update") into a project pre-seeded with the
    same keyword, so "update" has something to update. Returns (ImportResult, reloaded Keyword)."""
    payload = write_table(columns, rows, fmt)
    if fmt == "json":
        preview_rows = json.loads(payload)
        mapping = guess_mapping(columns)
    else:
        preview = read_table(f"export.{fmt}", payload)
        preview_rows = preview.rows
        mapping = guess_mapping(preview.columns)

    with session_factory() as s:
        fresh = Project(name="Fresh")
        s.add(fresh)
        s.commit()
        s.add(Keyword(project_id=fresh.id, keyword="buy server", source="manual"))
        s.commit()
        result = import_rows(s, fresh.id, preview_rows, mapping, "update", "test")
        s.commit()
        kw = s.scalars(select(Keyword).where(Keyword.project_id == fresh.id)).one()
        return result, kw


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
def test_export_reimport_round_trip_preserves_types_and_fixed_columns(session_factory, project, fmt):
    with session_factory() as s:
        s.add(
            Keyword(
                project_id=project.id,
                keyword="buy server",
                source="ahrefs",
                volume=100,
                kd=45.5,
                cpc=2.3,
                data=dict(ORIGINAL_DATA),
            )
        )
        s.commit()
        columns, rows = keyword_table(s, project.id, "all")
    assert "data.volume" in columns  # the collision was escaped, not silently dropped

    result, kw = _reimport_into_fresh_project(session_factory, columns, rows, fmt)

    assert result.updated == 1
    assert kw.volume == 100  # the fixed column, not the colliding data key (999)
    assert kw.kd == 45.5
    assert kw.cpc == 2.3
    assert kw.source == "ahrefs"
    assert kw.data == ORIGINAL_DATA


# --- formula-injection protection -------------------------------------------------------------


_FORMULA_ROW_COLUMNS = ["keyword", "volume", "kd", "cpc", "source", "cluster", "status", "reason"]
_FORMULA_ROW = {
    "keyword": "=1+2",
    "volume": None,
    "kd": None,
    "cpc": None,
    "source": "manual",
    "cluster": "",
    "status": "passing",
    "reason": "",
}


def test_write_csv_escapes_formula_injection():
    text = write_table(_FORMULA_ROW_COLUMNS, [_FORMULA_ROW], "csv").decode("utf-8-sig")
    assert "'=1+2" in text
    assert "\n=1+2" not in text


def test_write_xlsx_forces_text_for_formula_injection():
    columns = _FORMULA_ROW_COLUMNS
    payload = write_table(columns, [_FORMULA_ROW], "xlsx")
    wb = load_workbook(io.BytesIO(payload))
    cell = wb.active.cell(row=2, column=columns.index("keyword") + 1)
    assert cell.value == "=1+2"
    assert cell.data_type == "s"


def test_keyword_named_formula_is_escaped_on_export(session_factory, project):
    with session_factory() as s:
        s.add(Keyword(project_id=project.id, keyword="=1+2", source="manual"))
        s.commit()
        columns, rows = keyword_table(s, project.id, "all")

    csv_text = write_table(columns, rows, "csv").decode("utf-8-sig")
    assert "'=1+2" in csv_text

    wb = load_workbook(io.BytesIO(write_table(columns, rows, "xlsx")))
    cell = wb.active.cell(row=2, column=columns.index("keyword") + 1)
    assert cell.value == "=1+2"
    assert cell.data_type == "s"


FORMULA_PREFIXED_KEYWORDS = ["@home depot", "-50 sale", "+1 phone", "=cmd"]
FORMULA_PREFIXED_DATA = {"note": "-bad", "flag": "@risky"}


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_round_trip_with_formula_prefixed_keywords_and_data(session_factory, project, fmt):
    # Regression: _write_csv()'s leading "'" escape must be reversed on read, or reimporting
    # sito's own unedited export adds these as brand-new keywords instead of matching them.
    with session_factory() as s:
        for text in FORMULA_PREFIXED_KEYWORDS:
            s.add(Keyword(project_id=project.id, keyword=text, source="manual", data=dict(FORMULA_PREFIXED_DATA)))
        s.commit()
        columns, rows = keyword_table(s, project.id, "all")

    payload = write_table(columns, rows, fmt)
    preview = read_table(f"export.{fmt}", payload)
    mapping = guess_mapping(preview.columns)

    with session_factory() as s:
        fresh = Project(name="Fresh")
        s.add(fresh)
        s.commit()
        for text in FORMULA_PREFIXED_KEYWORDS:
            s.add(Keyword(project_id=fresh.id, keyword=text, source="manual"))
        s.commit()
        result = import_rows(s, fresh.id, preview.rows, mapping, "update", "test")
        s.commit()
        keywords = {k.keyword: k for k in s.scalars(select(Keyword).where(Keyword.project_id == fresh.id))}

    assert result.added == 0
    assert result.updated == len(FORMULA_PREFIXED_KEYWORDS)
    for text in FORMULA_PREFIXED_KEYWORDS:
        assert keywords[text].keyword == text
        assert keywords[text].data == FORMULA_PREFIXED_DATA


# --- parse_data_value / format_data_value are exact inverses --------------------------------


def test_parse_data_value_rejects_leading_zero_numbers():
    # "01234" (e.g. a zip code) must never be misread as the int 1234.
    assert parse_data_value("01234") == "01234"
    assert parse_data_value("0") == 0  # a lone "0" is still a valid int
    assert parse_data_value("0.5") == 0.5
    assert parse_data_value("01.5") == "01.5"


def test_format_data_value_quotes_strings_that_look_like_another_type():
    assert format_data_value("true") == '"true"'
    assert format_data_value("null") == '"null"'
    assert format_data_value("8") == '"8"'
    assert format_data_value('["a"]') == '"[\\"a\\"]"'
    assert format_data_value("01234") == "01234"  # already unambiguous, no quoting needed
    assert format_data_value("plain text") == "plain text"


def test_format_and_parse_data_value_are_exact_inverses():
    values = [8, 8.5, True, False, None, "01234", "true", "null", "plain text", ["a", "b"], {"h1": "Buy"}, '"quoted"']
    for value in values:
        assert parse_data_value(format_data_value(value)) == value
        if isinstance(value, str):
            assert isinstance(parse_data_value(format_data_value(value)), str)
