"""Import and export of the keyword table: file parsing, column mapping, writing.

Reading is deliberately forgiving (BOM sniffing, delimiter sniffing, several
encodings) because keyword lists come from wherever the user's tool of choice
exports them. Writing round-trips: a file exported by sito, edited in Excel
and re-imported with mode ``merge`` should update exactly what changed.
"""

from __future__ import annotations

import csv
import io
import json
import re
from collections.abc import Mapping
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from sito.core.snapshots import load_payload
from sito.core.text import normalize_keyword, to_float, to_int
from sito.models import Cluster, Keyword, Snapshot
from sito.plugins.base import PluginError

MEDIA_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
}

# --- reading -------------------------------------------------------------------------------


@dataclass
class TablePreview:
    columns: list[str]
    rows: list[dict[str, str]]


KEYWORD_HEADERS = {
    "keyword",
    "keywords",
    "key phrase",
    "search term",
    "query",
    "term",
    "ключевое слово",
    "ключевые слова",
    "запрос",
    "ключове слово",
    "ключові слова",
    "фраза",
}
VOLUME_HEADERS = {
    "volume",
    "search volume",
    "avg monthly searches",
    "avg. monthly searches",
    "monthly searches",
    "частотность",
    "частота",
    "объем",
    "объём",
    "обсяг",
}
KD_HEADERS = {"kd", "keyword difficulty", "difficulty", "keyword difficulty index"}
CPC_HEADERS = {"cpc", "cost per click"}
SOURCE_HEADERS = {"source", "источник", "джерело"}
STATUS_HEADERS = {"status"}
EXCLUDED_HEADERS = {"excluded"}
REASON_HEADERS = {"reason", "excluded_reason", "excluded reason"}
IGNORE_HEADERS = {"cluster"}

# The fixed columns keyword_table()/snapshot_table() always emit. A data key that collides
# with one of these is exported as "data.<key>" instead (see _table_rows) so it never
# overwrites or duplicates the fixed column; guess_mapping() reverses that below.
FIXED_COLUMNS = ("keyword", "volume", "kd", "cpc", "source", "cluster", "status", "reason")
DATA_COLUMN_PREFIX = "data."

# Cell text starting with any of these is a spreadsheet-formula-injection risk: a keyword,
# reason or custom field pasted verbatim into Excel/Sheets could run as a formula there.
# _write_csv() prefixes such text with a single quote; _read_delimited() reverses that.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _unescape_csv_cell(text: str) -> str:
    """Reverse the single-quote formula-injection escape _write_csv() adds."""
    if len(text) >= 2 and text[0] == "'" and text[1] in _FORMULA_PREFIXES:
        return text[1:]
    return text


def _decode(data: bytes) -> str:
    if data.startswith(b"\xff\xfe"):
        return data[2:].decode("utf-16-le")
    if data.startswith(b"\xfe\xff"):
        return data[2:].decode("utf-16-be")
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return data.decode("cp1251")
        except UnicodeDecodeError as exc:
            raise PluginError("Could not read this file as text (tried UTF-8, UTF-16 and CP1251).") from exc


def _norm_header(name: str) -> str:
    text = re.sub(r"\(.*?\)", "", name).strip().lower()
    return re.sub(r"\s+", " ", text)


def _sniff_delimiter(lines: list[str]) -> str | None:
    try:
        return csv.Sniffer().sniff("\n".join(lines), delimiters=",;\t|").delimiter
    except csv.Error:
        return None


def _dedupe_columns(columns: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for col in columns:
        name = col or "column"
        if name in seen:
            seen[name] += 1
            result.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 0
            result.append(name)
    return result


def _single_column_table(lines: list[str]) -> TablePreview:
    if _norm_header(lines[0]) in KEYWORD_HEADERS:
        lines = lines[1:]
    return TablePreview(columns=["Keyword"], rows=[{"Keyword": _unescape_csv_cell(line.strip())} for line in lines])


def _read_delimited(ext: str, data: bytes) -> TablePreview:
    lines = [line for line in _decode(data).splitlines() if line.strip()]
    if not lines:
        return TablePreview(columns=[], rows=[])
    if ext == ".txt":
        return _single_column_table(lines)
    delimiter = _sniff_delimiter(lines[:20])
    if delimiter is None:
        return _single_column_table(lines)

    rows_raw = list(csv.reader(lines, delimiter=delimiter))
    header = _dedupe_columns([cell.strip() or f"column_{i + 1}" for i, cell in enumerate(rows_raw[0])])
    rows: list[dict[str, str]] = []
    for raw in rows_raw[1:]:
        if not any(cell.strip() for cell in raw):
            continue
        raw = raw + [""] * (len(header) - len(raw))
        cells = (_unescape_csv_cell(cell.strip()) for cell in raw)
        rows.append(dict(zip(header, cells, strict=False)))
    return TablePreview(columns=header, rows=rows)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_xlsx(data: bytes) -> TablePreview:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises several exception types
        raise PluginError(f"Could not read this Excel file: {exc}") from exc
    try:
        rows_iter = wb.worksheets[0].iter_rows(values_only=True)
        try:
            header_cells = next(rows_iter)
        except StopIteration:
            return TablePreview(columns=[], rows=[])
        columns = _dedupe_columns([_cell_text(c) or f"column_{i + 1}" for i, c in enumerate(header_cells)])
        rows: list[dict[str, str]] = []
        for raw in rows_iter:
            values = [_cell_text(c) for c in raw]
            if not any(values):
                continue
            values += [""] * (len(columns) - len(values))
            rows.append(dict(zip(columns, values, strict=False)))
        return TablePreview(columns=columns, rows=rows)
    finally:
        wb.close()


def read_table(filename: str, data: bytes) -> TablePreview:
    ext = Path(filename).suffix.lower()
    if ext == ".xlsx":
        return _read_xlsx(data)
    if ext in (".csv", ".tsv", ".txt", ""):
        return _read_delimited(ext, data)
    raise PluginError(f"Unsupported file type '{ext or filename}'. Use .csv, .tsv, .txt or .xlsx.")


def slugify(name: str) -> str:
    text = re.sub(r"[^\w]+", "_", name.strip().lower()).strip("_")
    return text or "field"


def guess_mapping(columns: list[str]) -> dict[str, str]:
    """Column name -> target (``keyword``, ``volume``, ``data:<name>``, ...)."""
    mapping: dict[str, str] = {}
    for col in columns:
        if col.startswith(DATA_COLUMN_PREFIX):
            # A fixed-column-name collision, escaped on export as "data.<key>" - reverse it
            # literally rather than falling through to the header heuristics below.
            mapping[col] = f"data:{col[len(DATA_COLUMN_PREFIX) :]}"
            continue
        norm = _norm_header(col)
        if norm in KEYWORD_HEADERS:
            mapping[col] = "keyword"
        elif norm in VOLUME_HEADERS:
            mapping[col] = "volume"
        elif norm in KD_HEADERS:
            mapping[col] = "kd"
        elif norm in CPC_HEADERS:
            mapping[col] = "cpc"
        elif norm in SOURCE_HEADERS:
            mapping[col] = "source"
        elif norm in STATUS_HEADERS or norm in EXCLUDED_HEADERS:
            mapping[col] = "excluded"
        elif norm in REASON_HEADERS:
            mapping[col] = "excluded_reason"
        elif norm in IGNORE_HEADERS:
            mapping[col] = "ignore"
        else:
            mapping[col] = f"data:{slugify(col)}"
    return mapping


# --- importing -----------------------------------------------------------------------------


@dataclass
class ImportResult:
    added: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)


def _parse_bool_status(value: str) -> bool | None:
    v = value.strip().lower()
    if v in ("retained", "excluded", "true", "1", "yes", "y"):
        return True
    if v in ("passing", "active", "false", "0", "no", "n"):
        return False
    return None


# No leading zeros (so "01234", a zip code, is never mistaken for the int 1234).
_INT_RE = re.compile(r"^-?(0|[1-9]\d*)$")
_FLOAT_RE = re.compile(r"^-?(0|[1-9]\d*)\.\d+$")
_BOOL_NULL_WORDS = ("true", "false", "null", "none")


def parse_data_value(raw: str) -> Any:
    """Parse a data cell back to its likely original type - the exact inverse of format_data_value().

    A JSON-quoted string (``"..."``) decodes back to plain text, taking priority over every
    other rule below (that's how format_data_value() protects a string that would otherwise be
    misread). Otherwise: lists/dicts are JSON (``[...]``/``{...}``); ``true``/``false``/``null``/
    ``none`` become bool/None; plain integers/decimals (no leading zeros) become int/float;
    anything else stays text.
    """
    text = raw.strip()
    if not text:
        return raw
    if len(text) >= 2 and text[0] == '"' and text[-1] == '"':
        try:
            decoded = json.loads(text)
        except json.JSONDecodeError:
            decoded = None
        if isinstance(decoded, str):
            return decoded
    if text[0] in "[{":
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            return raw
    low = text.lower()
    if low == "true":
        return True
    if low == "false":
        return False
    if low in ("null", "none"):
        return None
    if _INT_RE.fullmatch(text):
        return int(text)
    if _FLOAT_RE.fullmatch(text):
        return float(text)
    return raw


def _looks_like_another_type(text: str) -> bool:
    """Whether parse_data_value(text) would return something other than the plain string ``text``."""
    if not text:
        return False
    if text[0] in '[{"':
        return True
    if text.lower() in _BOOL_NULL_WORDS:
        return True
    return bool(_INT_RE.fullmatch(text) or _FLOAT_RE.fullmatch(text))


def format_data_value(value: Any) -> str:
    """The inverse of :func:`parse_data_value`, for showing a data value as editable text.

    A string that would otherwise round-trip as a different type (numeric-looking, a
    true/false/null keyword, or itself starting with ``[``/``{``/``"``) is JSON-quoted so it
    comes back as the same string, not a bool/number/None/nested structure.
    """
    if isinstance(value, str):
        return json.dumps(value, ensure_ascii=False) if _looks_like_another_type(value) else value
    return json.dumps(value, ensure_ascii=False)


def _collect_row_values(row: Mapping[str, Any], mapping: dict[str, str]) -> tuple[dict[str, Any], dict[str, Any]]:
    values: dict[str, Any] = {}
    data: dict[str, Any] = {}
    for column, target in mapping.items():
        if target == "ignore":
            continue
        raw = row.get(column)
        if raw is None:
            continue
        if target.startswith("data:"):
            if isinstance(raw, str):
                text = raw.strip()
                if text != "":
                    data[target[5:]] = parse_data_value(text)
            else:
                data[target[5:]] = raw  # already a proper type, e.g. re-importing a JSON export
            continue
        raw = str(raw).strip()
        if target == "keyword":
            values["keyword"] = raw
        elif target == "excluded":
            parsed = _parse_bool_status(raw)
            if parsed is not None:
                values["excluded"] = parsed
        elif raw == "":
            continue  # empty values never overwrite volume/kd/cpc/source/excluded_reason
        elif target == "volume":
            values["volume"] = to_int(raw)
        elif target == "kd":
            values["kd"] = to_float(raw)
        elif target == "cpc":
            values["cpc"] = to_float(raw)
        elif target == "source":
            values["source"] = raw[:100]
        elif target == "excluded_reason":
            values["excluded_reason"] = raw[:300]
    return values, data


def import_rows(
    session: Session,
    project_id: int,
    rows: list[dict[str, str]],
    mapping: dict[str, str],
    mode: str,
    source_label: str,
) -> ImportResult:
    result = ImportResult()
    if "keyword" not in mapping.values():
        result.errors.append("No column is mapped to 'Keyword'; nothing was imported.")
        result.skipped = len(rows)
        return result

    existing = {
        normalize_keyword(k.keyword): k
        for k in session.scalars(select(Keyword).where(Keyword.project_id == project_id))
    }

    for index, row in enumerate(rows, start=1):
        values, data = _collect_row_values(row, mapping)
        norm = normalize_keyword(values.get("keyword"))
        if not norm or len(norm) > 500:
            result.skipped += 1
            if len(result.errors) < 20:
                result.errors.append(f"Row {index}: empty or too long keyword, skipped.")
            continue

        existing_kw = existing.get(norm)
        if existing_kw is None:
            if mode == "update":
                result.skipped += 1
                continue
            kw = Keyword(
                project_id=project_id,
                keyword=norm,
                source=(values.get("source") or source_label)[:100],
                volume=values.get("volume"),
                kd=values.get("kd"),
                cpc=values.get("cpc"),
                excluded=values.get("excluded", False),
                excluded_reason=values.get("excluded_reason"),
                data=data,
            )
            session.add(kw)
            existing[norm] = kw
            result.added += 1
        else:
            if mode == "add":
                result.skipped += 1
                continue
            for name in ("volume", "kd", "cpc", "source", "excluded", "excluded_reason"):
                if name in values:
                    setattr(existing_kw, name, values[name])
            if data:
                existing_kw.data = {**(existing_kw.data or {}), **data}
            result.updated += 1

    session.flush()
    return result


# --- exporting -----------------------------------------------------------------------------


def _cluster_name_map(session: Session, cluster_ids: set[int]) -> dict[int, str]:
    if not cluster_ids:
        return {}
    return dict(session.execute(select(Cluster.id, Cluster.name).where(Cluster.id.in_(cluster_ids))).all())


def _get(item: Any, name: str) -> Any:
    return item.get(name) if isinstance(item, Mapping) else getattr(item, name)


def _table_rows(items: list[Any], cluster_names: dict[int, str]) -> tuple[list[str], list[dict[str, Any]]]:
    data_keys: set[str] = set()
    prepared = []
    for item in items:
        data = _get(item, "data") or {}
        data_keys.update(data.keys())
        cluster_id = _get(item, "cluster_id")
        prepared.append(
            {
                "keyword": _get(item, "keyword"),
                "volume": _get(item, "volume"),
                "kd": _get(item, "kd"),
                "cpc": _get(item, "cpc"),
                "source": _get(item, "source"),
                "cluster": cluster_names.get(cluster_id, "") if cluster_id else "",
                "status": "excluded" if _get(item, "excluded") else "passing",
                "reason": _get(item, "excluded_reason") or "",
                "_data": data,
            }
        )
    sorted_keys = sorted(data_keys)
    # A data key named e.g. "volume" would otherwise silently overwrite/duplicate the fixed
    # "volume" column (same dict key); escape it as "data.volume" instead.
    data_columns = [f"{DATA_COLUMN_PREFIX}{key}" if key in FIXED_COLUMNS else key for key in sorted_keys]
    columns = [*FIXED_COLUMNS, *data_columns]
    rows = []
    for row in prepared:
        data = row.pop("_data")
        for key, col in zip(sorted_keys, data_columns, strict=True):
            row[col] = data.get(key, "")
        rows.append(row)
    return columns, rows


def keyword_table(session: Session, project_id: int, scope: str) -> tuple[list[str], list[dict[str, Any]]]:
    query = select(Keyword).where(Keyword.project_id == project_id)
    if scope == "active":
        query = query.where(Keyword.excluded.is_(False))
    keywords = list(session.scalars(query.order_by(Keyword.id)))
    cluster_names = _cluster_name_map(session, {k.cluster_id for k in keywords if k.cluster_id})
    return _table_rows(keywords, cluster_names)


def snapshot_table(snapshot: Snapshot, scope: str = "all") -> tuple[list[str], list[dict[str, Any]]]:
    payload = load_payload(snapshot)
    keywords = payload.get("keywords", [])
    if scope == "active":
        keywords = [k for k in keywords if not k.get("excluded")]
    cluster_names = {c["id"]: c["name"] for c in payload.get("clusters", [])}
    return _table_rows(keywords, cluster_names)


def _export_value(col: str, value: Any) -> Any:
    """The value to hand to the CSV/XLSX writer for one cell (before formula-escaping).

    Fixed columns keep their native type (str/number/None) untouched - they already have their
    own dedicated parsing on the way back in. Data columns go through format_data_value() so a
    stored type (bool/int/float/None/list/dict) - or a string that would otherwise be misread as
    one - round-trips exactly.
    """
    if col in FIXED_COLUMNS or value is None or value == "":
        return value
    return format_data_value(value)


def _csv_cell(col: str, value: Any) -> Any:
    value = _export_value(col, value)
    if value is None:
        return ""
    if isinstance(value, str) and value[:1] in _FORMULA_PREFIXES:
        return "'" + value
    return value


def _write_csv(columns: list[str], rows: list[dict[str, Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([_csv_cell(c, row.get(c)) for c in columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _write_xlsx(columns: list[str], rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for row in rows:
        values = [_export_value(c, row.get(c)) for c in columns]
        ws.append(values)
        for cell, value in zip(ws[ws.max_row], values, strict=True):
            if isinstance(value, str) and value[:1] in _FORMULA_PREFIXES:
                cell.data_type = "s"  # force text: don't let Excel treat it as a formula
    ws.auto_filter.ref = ws.dimensions
    for idx, name in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        width = len(str(name))
        for cell in ws[letter][1:200]:
            if cell.value is not None:
                width = max(width, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(width + 2, 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_json(columns: list[str], rows: list[dict[str, Any]]) -> bytes:
    payload = [{c: row.get(c) for c in columns} for row in rows]
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def write_table(columns: list[str], rows: list[dict[str, Any]], fmt: str) -> bytes:
    if fmt == "csv":
        return _write_csv(columns, rows)
    if fmt == "xlsx":
        return _write_xlsx(columns, rows)
    if fmt == "json":
        return _write_json(columns, rows)
    raise PluginError(f"Unsupported export format '{fmt}'.")


_SLUG_RE = re.compile(r"[^a-z0-9]+")


def export_filename(project_name: str, label: str, fmt: str) -> str:
    def slug(text: str) -> str:
        return _SLUG_RE.sub("-", text.strip().lower()).strip("-") or "export"

    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    return f"{slug(project_name)}-{slug(label)}-{stamp}.{fmt}"
